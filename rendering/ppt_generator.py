"""
PPT 生成器
使用 python-pptx 将阅读笔记/摘要/金句生成为 PowerPoint 演示文稿
"""
import logging
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import List, Optional

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

logger = logging.getLogger(__name__)


@dataclass
class SlideContent:
    """单页幻灯片内容"""
    layout: str = "quote"  # "title" | "quote" | "bullets" | "image_text"
    title: str = ""
    body: str = ""
    bullets: List[str] = field(default_factory=list)
    image_path: str = ""
    source: str = ""  # 来源（书名/作者）


@dataclass
class PPTTheme:
    """PPT 配色主题"""
    bg_color: tuple = (250, 250, 250)
    title_color: tuple = (26, 26, 26)
    body_color: tuple = (68, 68, 68)
    accent_color: tuple = (33, 150, 243)


THEMES = {
    "light": PPTTheme(),
    "warm": PPTTheme(
        bg_color=(253, 246, 236),
        title_color=(93, 64, 55),
        body_color=(109, 76, 65),
        accent_color=(121, 85, 72),
    ),
    "dark": PPTTheme(
        bg_color=(38, 38, 38),
        title_color=(255, 255, 255),
        body_color=(200, 200, 200),
        accent_color=(100, 181, 246),
    ),
}


class PPTGenerator:
    """PPT 演示文稿生成器"""

    def __init__(self, output_dir: str = "./data/ppt", theme: str = "light"):
        if not HAS_PPTX:
            raise ImportError("python-pptx 未安装，请运行: pip install python-pptx")
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.theme = THEMES.get(theme, THEMES["light"])

    def generate(
        self,
        slides: List[SlideContent],
        title: str = "阅读笔记",
        subtitle: str = "",
        filename: Optional[str] = None,
    ) -> str:
        """
        生成 PPT 文件，返回文件路径。

        Args:
            slides: 幻灯片内容列表
            title: 演示标题
            subtitle: 副标题
            filename: 文件名（不含扩展名）
        """
        prs = Presentation()
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        # 标题页
        self._add_title_slide(prs, title, subtitle or datetime.now().strftime("%Y年%m月%d日"))

        # 内容页
        for slide_data in slides:
            if slide_data.layout == "quote":
                self._add_quote_slide(prs, slide_data)
            elif slide_data.layout == "bullets":
                self._add_bullets_slide(prs, slide_data)
            elif slide_data.layout == "image_text":
                self._add_image_text_slide(prs, slide_data)
            else:
                self._add_quote_slide(prs, slide_data)

        fname = filename or f"reading_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        filepath = self.output_dir / f"{fname}.pptx"
        prs.save(str(filepath))
        logger.info(f"PPT 已生成: {filepath}")
        return str(filepath)

    def generate_from_notes(
        self,
        notes: List[dict],
        book_title: str = "",
        filename: Optional[str] = None,
    ) -> str:
        """
        从笔记列表快速生成 PPT。

        Args:
            notes: [{"content": "...", "tags": [...], "book_name": "..."}]
            book_title: 总标题用书名
            filename: 文件名
        """
        slides = []
        for note in notes:
            content = note.get("content", "")
            tags = note.get("tags", [])
            book = note.get("book_name", book_title)
            source = f"《{book}》" if book else ""

            if len(content) < 100:
                slides.append(SlideContent(
                    layout="quote", body=content, source=source,
                ))
            else:
                # 长内容分 bullets
                lines = [l.strip() for l in content.split("\n") if l.strip()]
                if len(lines) > 1:
                    slides.append(SlideContent(
                        layout="bullets",
                        title=tags[0] if tags else "笔记",
                        bullets=lines,
                        source=source,
                    ))
                else:
                    slides.append(SlideContent(
                        layout="quote", body=content, source=source,
                    ))

        title = f"《{book_title}》读书笔记" if book_title else "读书笔记"
        return self.generate(slides, title=title, filename=filename)

    def _add_title_slide(self, prs, title: str, subtitle: str):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        self._set_bg(slide)

        # 标题
        txBox = slide.shapes.add_textbox(
            Inches(1), Inches(2.5), Inches(11.33), Inches(1.5)
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(44)
        p.font.color.rgb = RGBColor(*self.theme.title_color)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

        # 副标题
        txBox2 = slide.shapes.add_textbox(
            Inches(1), Inches(4.2), Inches(11.33), Inches(1)
        )
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(24)
        p2.font.color.rgb = RGBColor(*self.theme.body_color)
        p2.alignment = PP_ALIGN.CENTER

    def _add_quote_slide(self, prs, data: SlideContent):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._set_bg(slide)

        # 引号装饰
        deco = slide.shapes.add_textbox(
            Inches(1), Inches(1.5), Inches(1), Inches(1)
        )
        dp = deco.text_frame.paragraphs[0]
        dp.text = "\u201C"
        dp.font.size = Pt(80)
        dp.font.color.rgb = RGBColor(*self.theme.accent_color)

        # 正文
        txBox = slide.shapes.add_textbox(
            Inches(1.5), Inches(2.5), Inches(10), Inches(3)
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = data.body
        p.font.size = Pt(28)
        p.font.color.rgb = RGBColor(*self.theme.body_color)
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(12)

        # 来源
        if data.source:
            txBox3 = slide.shapes.add_textbox(
                Inches(1), Inches(5.8), Inches(11.33), Inches(0.8)
            )
            tf3 = txBox3.text_frame
            p3 = tf3.paragraphs[0]
            p3.text = f"—— {data.source}"
            p3.font.size = Pt(18)
            p3.font.color.rgb = RGBColor(*self.theme.accent_color)
            p3.alignment = PP_ALIGN.RIGHT

    def _add_bullets_slide(self, prs, data: SlideContent):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._set_bg(slide)

        # 标题
        if data.title:
            txBox = slide.shapes.add_textbox(
                Inches(1), Inches(0.8), Inches(11.33), Inches(1)
            )
            tf = txBox.text_frame
            p = tf.paragraphs[0]
            p.text = data.title
            p.font.size = Pt(36)
            p.font.color.rgb = RGBColor(*self.theme.title_color)
            p.font.bold = True

        # Bullets
        txBox2 = slide.shapes.add_textbox(
            Inches(1.2), Inches(2), Inches(10.5), Inches(4.5)
        )
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        for i, bullet in enumerate(data.bullets):
            if i == 0:
                p = tf2.paragraphs[0]
            else:
                p = tf2.add_paragraph()
            p.text = f"  {bullet}"
            p.font.size = Pt(22)
            p.font.color.rgb = RGBColor(*self.theme.body_color)
            p.space_after = Pt(10)

            # bullet marker
            p.level = 0

        # 来源
        if data.source:
            txBox3 = slide.shapes.add_textbox(
                Inches(1), Inches(6.5), Inches(11.33), Inches(0.5)
            )
            tf3 = txBox3.text_frame
            p3 = tf3.paragraphs[0]
            p3.text = data.source
            p3.font.size = Pt(16)
            p3.font.color.rgb = RGBColor(*self.theme.accent_color)
            p3.alignment = PP_ALIGN.RIGHT

    def _add_image_text_slide(self, prs, data: SlideContent):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._set_bg(slide)

        # 左图
        if data.image_path and Path(data.image_path).exists():
            slide.shapes.add_picture(
                data.image_path,
                Inches(0.5), Inches(0.5),
                Inches(6), Inches(6.5),
            )
            text_left = Inches(7)
        else:
            text_left = Inches(1)

        text_width = Inches(12.33) - text_left

        # 标题
        if data.title:
            txBox = slide.shapes.add_textbox(
                text_left, Inches(1), text_width, Inches(1)
            )
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = data.title
            p.font.size = Pt(30)
            p.font.color.rgb = RGBColor(*self.theme.title_color)
            p.font.bold = True

        # 正文
        txBox2 = slide.shapes.add_textbox(
            text_left, Inches(2.2), text_width, Inches(4)
        )
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = data.body
        p2.font.size = Pt(20)
        p2.font.color.rgb = RGBColor(*self.theme.body_color)

    def _set_bg(self, slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*self.theme.bg_color)
